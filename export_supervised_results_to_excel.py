import argparse
import pickle
import re
from pathlib import Path

import numpy as np
import pandas as pd

from export_results_to_excel import (
    DEFAULT_AGGREGATORS,
    DEFAULT_CONFIDENCE,
    DEFAULT_EXPECTED_RUNS,
    DEFAULT_OUTPUT_DIR,
    DEFAULT_RESULTS_DIR,
    build_plot_tables,
    build_summary,
    safe_last,
    safe_max,
    safe_min,
    sanitize_sheet_name,
    scalar_or_nan,
)


DEFAULT_OUTPUT_NAME = "tgn_supervised_report.xlsx"
DEFAULT_DATASETS = [
    "toy_v3_last_event",
    "toy_v3_persistent_mean",
    "toy_v3_rare_spike",
    "toy_v3_ordered_pattern",
    "toy_last_event",
    "toy_persistent_mean",
    "toy_rare_spike",
    "toy_ordered_pattern",
]
DEFAULT_PREDICTION_TASK = "edge_label_classification"

PRIMARY_METRICS = [
    "test_ap",
    "test_auc",
    "test_f1",
]
SECONDARY_METRICS = [
    "test_acc",
    "test_prec",
    "test_rec",
    "best_val_auc",
    "final_val_auc",
    "best_val_ap",
    "final_val_ap",
    "best_val_f1",
    "final_val_f1",
    "final_train_loss",
    "min_train_loss",
    "avg_epoch_time",
    "total_train_time",
]
DEFAULT_METRICS = PRIMARY_METRICS + SECONDARY_METRICS


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Build CSV and Excel summaries from supervised TGN edge-label classification "
            "result pickle files."
        )
    )
    parser.add_argument("--results-dir", default=DEFAULT_RESULTS_DIR,
                        help="Directory containing .pkl files produced by train_supervised.py.")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR,
                        help="Directory for CSV, Excel, and optional PNG chart outputs.")
    parser.add_argument("--output-name", default=DEFAULT_OUTPUT_NAME,
                        help="Excel workbook file name.")
    parser.add_argument("--datasets", nargs="*", default=DEFAULT_DATASETS,
                        help="Dataset names to include. Used to parse filenames reliably.")
    parser.add_argument("--aggregators", nargs="*", default=DEFAULT_AGGREGATORS,
                        help="Aggregator names to include.")
    parser.add_argument("--prefixes", nargs="*", default=None,
                        help="Optional prefix filter. Leave unset to accept any training prefix.")
    parser.add_argument("--metrics", nargs="*", default=DEFAULT_METRICS,
                        help="Metrics to summarize.")
    parser.add_argument("--chart-metrics", nargs="*", default=PRIMARY_METRICS,
                        help="Metrics to make chart-ready Excel sheets and PNG charts for.")
    parser.add_argument("--error-bars", default="std", choices=["std", "ci", "sem", "none"],
                        help="Error bars used in embedded/PNG charts. Default: standard deviation.")
    parser.add_argument("--prediction-task", default=DEFAULT_PREDICTION_TASK,
                        help="Prediction task value expected inside result pickle files.")
    parser.add_argument("--include-any-task", action="store_true",
                        help="Include legacy supervised result files from any prediction task.")
    parser.add_argument("--confidence", type=float, default=DEFAULT_CONFIDENCE,
                        help="Confidence interval level, for example 0.95.")
    parser.add_argument("--expected-runs", type=int, default=DEFAULT_EXPECTED_RUNS,
                        help="Expected number of runs per dataset/aggregator pair.")
    parser.add_argument("--include-incomplete", action="store_true",
                        help="Include partially written pickle files without test metrics.")
    parser.add_argument("--skip-png-charts", action="store_true",
                        help="Skip PNG bar charts and embedded chart images. Excel/CSV tables are still written.")
    return parser.parse_args()


def parse_supervised_result_filename(path, datasets, aggregators):
    aggregator_pattern = "|".join(re.escape(a) for a in sorted(aggregators, key=len, reverse=True))
    match = re.match(
        rf"^supervised_(?P<body>.+)_(?P<aggregator>{aggregator_pattern})"
        rf"_(?P<result_task>node_classification|edge_label_classification)(?:_(?P<run>\d+))?$",
        path.stem,
    )
    if not match:
        return None

    body = match.group("body")
    dataset_candidates = [
        dataset for dataset in datasets
        if body == dataset or body.endswith(f"_{dataset}")
    ]
    if not dataset_candidates:
        return None

    dataset = max(dataset_candidates, key=len)
    prefix = body[:-len(dataset)]
    if prefix.endswith("_"):
        prefix = prefix[:-1]
    if prefix == "_":
        prefix = ""

    return {
        "dataset": dataset,
        "prefix": prefix,
        "aggregator": match.group("aggregator"),
        "result_task": match.group("result_task"),
        "run": int(match.group("run") or 0),
    }


def load_supervised_run_file(path):
    with open(path, "rb") as handle:
        data = pickle.load(handle)

    val_aucs = data.get("val_aucs", []) or []
    val_aps = data.get("val_aps", []) or []
    val_f1s = data.get("val_f1s", []) or []
    train_losses = data.get("train_losses", []) or []
    total_epoch_times = data.get("total_epoch_times", []) or []

    row = {
        "seed": data.get("seed", np.nan),
        "prediction_task": data.get("prediction_task", "legacy_node_classification"),
        "test_ap": scalar_or_nan(data.get("test_ap")),
        "test_auc": scalar_or_nan(data.get("test_auc")),
        "test_acc": scalar_or_nan(data.get("test_acc")),
        "test_prec": scalar_or_nan(data.get("test_prec")),
        "test_rec": scalar_or_nan(data.get("test_rec")),
        "test_f1": scalar_or_nan(data.get("test_f1")),
        "best_val_auc": safe_max(val_aucs),
        "final_val_auc": safe_last(val_aucs),
        "best_val_ap": safe_max(val_aps),
        "final_val_ap": safe_last(val_aps),
        "best_val_f1": safe_max(val_f1s),
        "final_val_f1": safe_last(val_f1s),
        "final_train_loss": safe_last(train_losses),
        "min_train_loss": safe_min(train_losses),
        "avg_epoch_time": float(np.mean(total_epoch_times)) if len(total_epoch_times) else np.nan,
        "total_train_time": float(np.sum(total_epoch_times)) if len(total_epoch_times) else np.nan,
        "num_epochs_recorded": len(total_epoch_times),
        "edge_decoder_input_dim": data.get("edge_decoder_input_dim", np.nan),
        "label_filter": data.get("label_filter", "all_edges"),
        "label_control": data.get("label_control", "none"),
        "decoder_input_control": data.get("decoder_input_control", "full"),
        "explicit_memory_replay": data.get("explicit_memory_replay", False),
        "test_n_eval": data.get("test_n_eval", np.nan),
        "val_n_eval": data.get("val_n_eval", np.nan),
        "source_encoder_model": data.get("source_encoder_model", ""),
        "decoder_model": data.get("decoder_model", ""),
        "source_file": path.name,
    }
    row["is_complete"] = not np.isnan(row["test_auc"])
    return row


def collect_raw_runs(results_dir, datasets, aggregators, prefixes=None, prediction_task=None,
                     include_any_task=False, include_incomplete=False):
    rows = []
    prefixes = set(prefixes or [])

    for path in sorted(results_dir.glob("supervised_*_classification*.pkl")):
        parsed = parse_supervised_result_filename(path, datasets=datasets, aggregators=aggregators)
        if parsed is None:
            continue
        if prefixes and parsed["prefix"] not in prefixes:
            continue

        row = load_supervised_run_file(path)
        if not include_any_task and prediction_task and row["prediction_task"] != prediction_task:
            continue
        if not include_incomplete and not row["is_complete"]:
            continue
        row.update(parsed)
        rows.append(row)

    if not rows:
        raise FileNotFoundError(
            f"No matching supervised result .pkl files found in {results_dir}. "
            "Check --datasets, --aggregators, --prefixes, and --prediction-task. "
            "If you want to include older node-only supervised files, pass --include-any-task."
        )

    return (
        pd.DataFrame(rows)
        .sort_values(["dataset", "aggregator", "run", "source_file"])
        .reset_index(drop=True)
    )


def build_completion_table(summary_long, expected_runs):
    test_auc = summary_long[summary_long["metric"] == "test_auc"].copy()
    if test_auc.empty:
        return pd.DataFrame()
    completion = test_auc[["dataset", "aggregator", "n_runs", "source_files"]].copy()
    completion["expected_runs"] = expected_runs
    completion["missing_runs"] = expected_runs - completion["n_runs"]
    completion["complete"] = completion["missing_runs"] <= 0
    return completion.sort_values(["dataset", "aggregator"]).reset_index(drop=True)


def write_plot_sheet(writer, metric, tables):
    sheet_name = sanitize_sheet_name(f"plot_{metric}")
    row = 0
    for label, table in tables.items():
        title_df = pd.DataFrame({f"{metric}_{label}": []})
        title_df.to_excel(writer, sheet_name=sheet_name, startrow=row, index=False)
        row += 1
        table.reset_index().to_excel(writer, sheet_name=sheet_name, startrow=row, index=False)
        row += len(table) + 4


def add_chart_images_to_workbook(workbook, chart_paths, error_bars):
    if not chart_paths:
        return False

    try:
        from openpyxl.drawing.image import Image
    except ImportError:
        return False

    sheet_name = sanitize_sheet_name(f"charts_{error_bars}")
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    worksheet["A1"] = f"Embedded bar charts with {error_bars} error bars"

    row = 3
    for chart_path in chart_paths:
        worksheet[f"A{row}"] = chart_path.stem
        image = Image(str(chart_path))
        image.width = 960
        image.height = 480
        worksheet.add_image(image, f"A{row + 1}")
        row += 29

    return True


def write_outputs(raw_df, summary_long, summary_wide, completion, plot_tables, output_dir,
                  output_name, confidence, expected_runs, chart_paths=None, error_bars="std"):
    output_dir.mkdir(parents=True, exist_ok=True)

    raw_csv = output_dir / "supervised_raw_runs.csv"
    summary_long_csv = output_dir / "supervised_summary_long.csv"
    summary_wide_csv = output_dir / "supervised_summary_wide.csv"
    completion_csv = output_dir / "supervised_run_completion.csv"

    raw_df.to_csv(raw_csv, index=False)
    summary_long.to_csv(summary_long_csv, index=False)
    summary_wide.to_csv(summary_wide_csv, index=False)
    completion.to_csv(completion_csv, index=False)

    workbook_path = output_dir / output_name
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        info = pd.DataFrame([
            {"field": "task", "value": DEFAULT_PREDICTION_TASK},
            {"field": "confidence_level", "value": confidence},
            {"field": "expected_runs", "value": expected_runs},
            {"field": "num_raw_rows", "value": len(raw_df)},
            {"field": "datasets", "value": ", ".join(sorted(raw_df["dataset"].unique()))},
            {"field": "aggregators", "value": ", ".join(sorted(raw_df["aggregator"].unique()))},
            {"field": "primary_metrics", "value": ", ".join(PRIMARY_METRICS)},
            {"field": "chart_error_bars", "value": error_bars},
            {
                "field": "metric_note",
                "value": (
                    "Current train_supervised.py predicts edge labels from source embedding, "
                    "destination embedding, and edge features. V3 toy runs use query_mask=1 "
                    "rows by default."
                ),
            },
        ])
        info.to_excel(writer, sheet_name="report_info", index=False)
        completion.to_excel(writer, sheet_name="run_completion", index=False)
        raw_df.to_excel(writer, sheet_name="raw_runs", index=False)
        summary_long.to_excel(writer, sheet_name="summary_long", index=False)
        summary_wide.to_excel(writer, sheet_name="summary_wide", index=False)
        for metric, tables in plot_tables.items():
            write_plot_sheet(writer, metric, tables)
        embedded_charts = add_chart_images_to_workbook(writer.book, chart_paths or [], error_bars)

    return {
        "workbook": workbook_path,
        "raw_csv": raw_csv,
        "summary_long_csv": summary_long_csv,
        "summary_wide_csv": summary_wide_csv,
        "completion_csv": completion_csv,
        "embedded_charts": embedded_charts,
    }


def get_error_table(tables, error_bars, confidence):
    means = tables["mean"]
    if error_bars == "none":
        return None, "none"
    if error_bars == "std":
        return tables["std"].reindex(index=means.index, columns=means.columns).fillna(0.0), "standard deviation"
    if error_bars == "sem":
        std = tables["std"].reindex(index=means.index, columns=means.columns)
        n_runs = tables["n_runs"].reindex(index=means.index, columns=means.columns)
        return (std / np.sqrt(n_runs)).fillna(0.0), "standard error"
    return (
        tables["ci_margin"].reindex(index=means.index, columns=means.columns).fillna(0.0),
        f"{int(confidence * 100)}% CI",
    )


def save_png_charts(plot_tables, output_dir, confidence, error_bars):
    try:
        import matplotlib.pyplot as plt
    except ImportError:
        return []

    charts_dir = output_dir / "supervised_charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    chart_paths = []

    for metric, tables in plot_tables.items():
        means = tables["mean"]
        errors, error_label = get_error_table(tables, error_bars, confidence)
        if means.empty:
            continue

        datasets = list(means.index)
        aggregators = list(means.columns)
        x = np.arange(len(datasets))
        width = 0.8 / max(len(aggregators), 1)

        fig, ax = plt.subplots(figsize=(12, 6))
        for idx, aggregator in enumerate(aggregators):
            offsets = x + (idx - (len(aggregators) - 1) / 2.0) * width
            ax.bar(
                offsets,
                means[aggregator].to_numpy(dtype=float),
                width=width,
                label=aggregator,
                yerr=None if errors is None else errors[aggregator].to_numpy(dtype=float),
                capsize=4,
            )

        ax.set_title(f"supervised {metric}: mean by dataset and aggregator")
        ax.set_xlabel("Dataset")
        ax.set_ylabel(metric)
        ax.set_xticks(x)
        ax.set_xticklabels(datasets, rotation=20, ha="right")
        ax.legend(title="Aggregator")
        ax.grid(axis="y", alpha=0.3)
        ax.text(
            0.99,
            0.01,
            f"Error bars: {error_label}",
            transform=ax.transAxes,
            ha="right",
            va="bottom",
            fontsize=9,
        )
        fig.tight_layout()

        chart_path = charts_dir / f"supervised_{metric}_bar_chart_{error_bars}.png"
        fig.savefig(chart_path, dpi=200, bbox_inches="tight")
        plt.close(fig)
        chart_paths.append(chart_path)

    return chart_paths


def print_compact_summary(summary_long, completion, paths, chart_paths):
    print("\nRun completion")
    if completion.empty:
        print("  No completion table available.")
    else:
        print(completion.to_string(index=False))

    print("\nPrimary metric summary")
    primary = summary_long[summary_long["metric"].isin(PRIMARY_METRICS)].copy()
    primary = primary[["dataset", "aggregator", "metric", "n_runs", "mean", "std", "ci_margin"]]
    with pd.option_context("display.max_rows", 200, "display.width", 140):
        print(primary.to_string(index=False, float_format=lambda value: f"{value:.4f}"))

    print("\nSaved outputs")
    for key, path in paths.items():
        print(f"  {key}: {path}")
    if chart_paths:
        print("  png_charts:")
        for chart_path in chart_paths:
            print(f"    {chart_path}")


def main():
    args = parse_args()
    results_dir = Path(args.results_dir)
    output_dir = Path(args.output_dir)

    raw_df = collect_raw_runs(
        results_dir=results_dir,
        datasets=args.datasets,
        aggregators=args.aggregators,
        prefixes=args.prefixes,
        prediction_task=args.prediction_task,
        include_any_task=args.include_any_task,
        include_incomplete=args.include_incomplete,
    )
    summary_long, summary_wide = build_summary(raw_df, args.metrics, args.confidence)
    completion = build_completion_table(summary_long, args.expected_runs)
    plot_tables = build_plot_tables(summary_long, args.chart_metrics)
    chart_paths = [] if args.skip_png_charts else save_png_charts(
        plot_tables,
        output_dir,
        args.confidence,
        args.error_bars,
    )
    paths = write_outputs(
        raw_df=raw_df,
        summary_long=summary_long,
        summary_wide=summary_wide,
        completion=completion,
        plot_tables=plot_tables,
        output_dir=output_dir,
        output_name=args.output_name,
        confidence=args.confidence,
        expected_runs=args.expected_runs,
        chart_paths=chart_paths,
        error_bars=args.error_bars,
    )
    print_compact_summary(summary_long, completion, paths, chart_paths)


if __name__ == "__main__":
    main()
